#!/usr/bin/env python3


from __future__ import annotations

import argparse
import ast
import math
from pathlib import Path
from typing import Dict, List, Set, Tuple, Optional

import pandas as pd
import networkx as nx

import matplotlib.pyplot as plt

from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image as XLImage


plt.rcParams["font.sans-serif"] = ["SimHei"]
plt.rcParams["axes.unicode_minus"] = False


TYPE_ALIASES = {
    "ISSUE": "问题陈述",
    "PHEN": "用户感知现象",
    "DIAG": "系统诊断信息",
    "MOD":  "功能模块",
    "OP":   "用户操作",
    "ELEM": "影响元素",
}

def normalize_type(t: str) -> str:
    if t is None:
        return ""
    s = str(t).strip()
    if not s or s.lower() == "nan":
        return ""
    return TYPE_ALIASES.get(s, s)


def parse_report_ids(x) -> Set[str]:


    if x is None or (isinstance(x, float) and math.isnan(x)):
        return set()
    if isinstance(x, (list, set, tuple)):
        return {str(i).strip() for i in x if str(i).strip()}
    s = str(x).strip()
    if not s or s.lower() in {"nan", "none", "null", "[]"}:
        return set()


    if (s.startswith("[") and s.endswith("]")) or (s.startswith("(") and s.endswith(")")):
        try:
            v = ast.literal_eval(s)
            if isinstance(v, (list, tuple, set)):
                return {str(i).strip() for i in v if str(i).strip()}
        except Exception:
            pass


    if "|" in s:
        parts = [p.strip() for p in s.split("|")]
        return {p for p in parts if p}


    for sep in [",", ";", " "]:
        if sep in s:
            parts = [p.strip() for p in s.split(sep)]
            return {p for p in parts if p}

    return {s}


def safe_read_excel(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"Not found: {path}")
    return pd.read_excel(path, dtype=str)


def build_graph(nodes_df: pd.DataFrame, edges_df: pd.DataFrame) -> Tuple[nx.Graph, Dict[str, Set[str]], Dict[str, str]]:

    for c in ["id", "name", "type", "source_row_index"]:
        if c not in nodes_df.columns:
            raise ValueError(f"nodes.xlsx missing column: {c}. Found: {list(nodes_df.columns)}")

    for c in ["source_id", "target_id", "relation", "source_row_index", "source_type", "target_type"]:
        if c not in edges_df.columns:
            raise ValueError(f"edges.xlsx missing column: {c}. Found: {list(edges_df.columns)}")

    G = nx.Graph()
    node_reports: Dict[str, Set[str]] = {}
    node_type: Dict[str, str] = {}


    for _, r in nodes_df.iterrows():
        nid = str(r["id"]).strip()
        if not nid or nid.lower() == "nan":
            continue
        nname = str(r["name"]).strip()
        ntype = normalize_type(r["type"])
        reps = parse_report_ids(r["source_row_index"])
        node_reports[nid] = reps
        node_type[nid] = ntype
        G.add_node(nid, entity_name=nname, entity_type=ntype, report_ids=reps)


    for _, r in edges_df.iterrows():
        u = str(r["source_id"]).strip()
        v = str(r["target_id"]).strip()
        if not u or not v or u.lower() == "nan" or v.lower() == "nan":
            continue

        if u not in G:
            G.add_node(u, entity_name=u, entity_type=normalize_type(r.get("source_type", "")), report_ids=set())
            node_reports[u] = set()
            node_type[u] = normalize_type(r.get("source_type", ""))
        if v not in G:
            G.add_node(v, entity_name=v, entity_type=normalize_type(r.get("target_type", "")), report_ids=set())
            node_reports[v] = set()
            node_type[v] = normalize_type(r.get("target_type", ""))

        rel = str(r["relation"]).strip()
        ereps = parse_report_ids(r["source_row_index"])

        if G.has_edge(u, v):
            prev = G[u][v].get("report_ids", set())
            G[u][v]["report_ids"] = set(prev) | set(ereps)
            rels = G[u][v].get("relation_types", set())
            if not isinstance(rels, set):
                rels = set([str(rels)])
            if rel:
                rels.add(rel)
            G[u][v]["relation_types"] = rels
        else:
            G.add_edge(u, v, report_ids=set(ereps), relation_types=set([rel]) if rel else set())

    return G, node_reports, node_type


def connectivity_metrics(G: nx.Graph) -> Tuple[int, int, float]:
    comps = list(nx.connected_components(G))
    num_cc = len(comps)
    if num_cc == 0:
        return 0, 0, 0.0
    lcc_size = max(len(c) for c in comps)
    lcc_ratio = lcc_size / max(1, G.number_of_nodes())
    return num_cc, lcc_size, lcc_ratio


def cross_report_edge_ratio(G: nx.Graph, node_reports: Dict[str, Set[str]]) -> Tuple[int, int, float]:
    total = G.number_of_edges()
    if total == 0:
        return 0, 0, 0.0
    cross = 0
    for u, v in G.edges():
        ru = node_reports.get(u, set())
        rv = node_reports.get(v, set())
        if (ru | rv) and (ru != rv):
            cross += 1
    return total, cross, cross / total


def node_in_universe(ntype: str, universe: Set[str]) -> bool:
    n = normalize_type(ntype)
    return bool(n) and (n in universe)


def candidate_subgraph_reports(
    G: nx.Graph,
    seed: str,
    khop: int,
    node_reports: Dict[str, Set[str]],
    node_type: Dict[str, str],
    universe: Set[str],
) -> Set[str]:
    if seed not in G:
        return set()

    nodes = {seed}
    frontier = {seed}
    for _ in range(khop):
        nxt = set()
        for x in frontier:
            nxt |= set(G.neighbors(x))
        nxt -= nodes
        nodes |= nxt
        frontier = nxt
        if not frontier:
            break

    covered: Set[str] = set()
    for n in nodes:
        if node_in_universe(node_type.get(n, ""), universe):
            covered |= node_reports.get(n, set())
    return covered


def build_candidates(
    G: nx.Graph,
    node_reports: Dict[str, Set[str]],
    node_type: Dict[str, str],
    universe: Set[str],
    khop: int,
    max_candidates: int,
) -> List[Tuple[str, Set[str]]]:
    scored = []
    for n in G.nodes():
        base = len(node_reports.get(n, set())) if node_in_universe(node_type.get(n, ""), universe) else 0
        score = base + math.log1p(G.degree(n))
        scored.append((score, n))
    scored.sort(reverse=True)
    seeds = [n for _, n in scored[:max_candidates]]

    candidates = []
    for s in seeds:
        cov = candidate_subgraph_reports(G, s, khop, node_reports, node_type, universe)
        if cov:
            candidates.append((s, cov))


    seen = set()
    uniq = []
    for seed, cov in candidates:
        key = frozenset(cov)
        if key in seen:
            continue
        seen.add(key)
        uniq.append((seed, cov))
    return uniq


def greedy_set_cover_curve(
    candidates: List[Tuple[str, Set[str]]],
    all_reports: Set[str],
    K: int,
) -> Tuple[List[float], float, List[str]]:
    uncovered = set(all_reports)
    chosen = []
    curve = []
    auc = 0.0
    remaining = candidates[:]

    for _k in range(1, K + 1):
        best_i = -1
        best_gain = -1
        for i, (_, cov) in enumerate(remaining):
            gain = len(cov & uncovered)
            if gain > best_gain:
                best_gain = gain
                best_i = i

        if best_i == -1 or best_gain <= 0:
            cov_ratio = 1.0 - (len(uncovered) / max(1, len(all_reports)))
            curve.append(cov_ratio)
            auc += cov_ratio
            continue

        seed, cov = remaining.pop(best_i)
        chosen.append(seed)
        uncovered -= (cov & uncovered)

        cov_ratio = 1.0 - (len(uncovered) / max(1, len(all_reports)))
        curve.append(cov_ratio)
        auc += cov_ratio

    return curve, auc, chosen


def load_aligned_clusters(path: Path) -> Optional[pd.DataFrame]:
    if not path.exists():
        return None
    df = pd.read_excel(path, dtype=str)

    if "cluster" not in df.columns or "entity" not in df.columns:
        return None
    df["cluster"] = df["cluster"].astype(str).str.strip()
    df["entity"] = df["entity"].astype(str).str.strip()
    df = df[(df["cluster"] != "") & (df["entity"] != "") & (df["cluster"].str.lower() != "nan") & (df["entity"].str.lower() != "nan")]
    return df


def compute_unification_stats(mapping_df: pd.DataFrame, topn: int = 20) -> Tuple[dict, pd.DataFrame]:
    m = mapping_df.copy()
    raw_entities = m["entity"].nunique(dropna=True)
    unified_entities = m["cluster"].nunique(dropna=True)
    reduction_ratio = 1.0 - (unified_entities / max(1, raw_entities))

    size_df = m.groupby("cluster")["entity"].nunique().reset_index(name="cluster_size").sort_values("cluster_size", ascending=False)
    avg_cluster_size = float(raw_entities / max(1, unified_entities))
    max_cluster_size = int(size_df["cluster_size"].max()) if len(size_df) else 0
    singleton_ratio = float((size_df["cluster_size"] == 1).mean()) if len(size_df) else 0.0

    top_clusters = size_df.head(topn)["cluster"].tolist()
    ex = (
        m[m["cluster"].isin(top_clusters)]
        .groupby("cluster")["entity"]
        .apply(lambda s: " | ".join(list(dict.fromkeys(s.tolist()))[:10]))
        .reset_index(name="example_entities")
    )
    top_df = size_df.head(topn).merge(ex, on="cluster", how="left")

    stats_row = dict(
        raw_entities=int(raw_entities),
        unified_entities=int(unified_entities),
        reduction_ratio=float(reduction_ratio),
        avg_cluster_size=float(avg_cluster_size),
        max_cluster_size=int(max_cluster_size),
        singleton_ratio=float(singleton_ratio),
    )
    return stats_row, top_df


def write_df(ws, df: pd.DataFrame):
    ws.append(list(df.columns))
    for r in dataframe_to_rows(df, index=False, header=False):
        ws.append(list(r))


def autosize_columns(ws, max_width=70):
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max_width, max(10, max_len + 2))


def plot_coverage_per_app(df_curve: pd.DataFrame, out_png: Path, app_id: str):
    plt.figure()
    sub = df_curve[df_curve["app_id"] == app_id].copy()
    for variant in ["no_unify", "unify"]:
        s = sub[sub["variant"] == variant].sort_values("k")
        plt.plot(s["k"], s["report_coverage"], label=variant)
    plt.xlabel("k (subgraph budget)")
    plt.ylabel("ReportCoverage@k")
    plt.title(f"Coverage Curve – {app_id}")
    plt.legend()
    plt.tight_layout()
    plt.savefig(out_png, dpi=200)
    plt.close()


def plot_coverage_across_apps(df_curve: pd.DataFrame, out_png: Path):
    plt.figure()
    g = df_curve.groupby(["variant", "k"])["report_coverage"]
    mean = g.mean().reset_index(name="mean")
    std = g.std(ddof=0).reset_index(name="std")
    merged = mean.merge(std, on=["variant", "k"], how="left").fillna({"std": 0.0})

    for variant in ["no_unify", "unify"]:
        s = merged[merged["variant"] == variant].sort_values("k")
        plt.plot(s["k"], s["mean"], label=f"{variant} (mean)")
        plt.fill_between(s["k"], s["mean"] - s["std"], s["mean"] + s["std"], alpha=0.2)

    plt.xlabel("k (subgraph budget)")
    plt.ylabel("ReportCoverage@k")
    plt.title("Coverage Curve – Across Apps (mean ± std)")
    plt.legend()
    plt.tight_layout()
    plt.savefig(out_png, dpi=200)
    plt.close()


def load_app_variant(root: Path, app_id: str, variant: str) -> Tuple[pd.DataFrame, pd.DataFrame]:
    npath = root / app_id / variant / "nodes.xlsx"
    epath = root / app_id / variant / "edges.xlsx"
    return safe_read_excel(npath), safe_read_excel(epath)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True, type=str, help="ROOT_glm4air folder")
    ap.add_argument("--out", default="RQ1_entity_unification_results.xlsx", type=str, help="Output xlsx path")
    ap.add_argument("--K", default=50, type=int, help="Budget K for coverage curve (1..K)")
    ap.add_argument("--khop", default=2, type=int, help="k-hop neighborhood for candidates")
    ap.add_argument("--universe", nargs="+", default=["ISSUE", "PHEN", "DIAG"], help="Universe types (ISSUE/PHEN/DIAG or Chinese)")
    ap.add_argument("--max_candidates", default=400, type=int, help="Max candidate seeds per app per variant")
    ap.add_argument("--seed_test_n", default=0, type=int, help="If >0: run seed expansion test with N seeds (from unify)")
    args = ap.parse_args()

    root = Path(args.root)
    out_xlsx = Path(args.out)
    out_dir = out_xlsx.parent
    out_dir.mkdir(parents=True, exist_ok=True)

    universe = {normalize_type(x) for x in args.universe}


    apps = sorted([p.name for p in root.iterdir() if p.is_dir()])
    if not apps:
        raise RuntimeError(f"No app folders found under: {root}")


    rows_graph_stats = []
    rows_conn = []
    rows_cross = []
    rows_curve = []
    rows_covsum = []
    rows_delta = []
    rows_seed = []
    rows_unify_stats = []
    unify_top_tables = []


    readme = pd.DataFrame(
        [
            ["Experiment", "RQ1 – Entity Unification Ablation"],
            ["Description", "Structural & Coverage Impact of Entity Unification"],
            ["Variants", "no_unify, unify"],
            ["Graph Type", "Undirected Simple Graph (direction ignored, multi-edges merged)"],
            ["Coverage Algorithm", "Greedy Set Cover over k-hop neighborhoods"],
            ["Universe", ", ".join(sorted(universe))],
            ["Budget K", str(args.K)],
            ["k-hop", str(args.khop)],
            ["Notes", "All inputs share identical extraction; no gold required."],
        ],
        columns=["Field", "Value"],
    )

    for app_id in apps:

        try:
            nodes_no, edges_no = load_app_variant(root, app_id, "no_unify")
            nodes_u,  edges_u  = load_app_variant(root, app_id, "unify")
        except FileNotFoundError:

            continue


        mdf = load_aligned_clusters(root / app_id / "aligned_clusters.xlsx")
        if mdf is not None and len(mdf) > 0:
            stats_row, top_df = compute_unification_stats(mdf, topn=20)
            rows_unify_stats.append(dict(app_id=app_id, **stats_row))
            top_df.insert(0, "app_id", app_id)
            unify_top_tables.append(top_df)


        G_no, rep_no, type_no = build_graph(nodes_no, edges_no)
        G_u,  rep_u,  type_u  = build_graph(nodes_u,  edges_u)


        rows_graph_stats += [
            dict(app_id=app_id, variant="no_unify", num_nodes=G_no.number_of_nodes(), num_edges=G_no.number_of_edges()),
            dict(app_id=app_id, variant="unify",    num_nodes=G_u.number_of_nodes(),  num_edges=G_u.number_of_edges()),
        ]


        cc_no, lcc_size_no, lcc_ratio_no = connectivity_metrics(G_no)
        cc_u,  lcc_size_u,  lcc_ratio_u  = connectivity_metrics(G_u)
        rows_conn += [
            dict(app_id=app_id, variant="no_unify", num_cc=cc_no, lcc_size=lcc_size_no, lcc_ratio=lcc_ratio_no),
            dict(app_id=app_id, variant="unify",    num_cc=cc_u,  lcc_size=lcc_size_u,  lcc_ratio=lcc_ratio_u),
        ]


        tot_no, cross_no, cer_no = cross_report_edge_ratio(G_no, rep_no)
        tot_u,  cross_u,  cer_u  = cross_report_edge_ratio(G_u,  rep_u)
        rows_cross += [
            dict(app_id=app_id, variant="no_unify", num_edges=tot_no, cross_report_edges=cross_no, cross_report_ratio=cer_no),
            dict(app_id=app_id, variant="unify",    num_edges=tot_u,  cross_report_edges=cross_u,  cross_report_ratio=cer_u),
        ]


        def run_variant(variant_name: str, G: nx.Graph, rep: Dict[str, Set[str]], typ: Dict[str, str]):

            all_reports = set()
            for rset in rep.values():
                all_reports |= set(rset)


            universe_nodes = sum(1 for n in G.nodes() if node_in_universe(typ.get(n, ""), universe))
            print(f"[{app_id} | {variant_name}] reports={len(all_reports)} universe_nodes={universe_nodes}")


            cands = build_candidates(G, rep, typ, universe, args.khop, args.max_candidates)
            print(f"[{app_id} | {variant_name}] candidates={len(cands)}")

            curve, auc, chosen = greedy_set_cover_curve(cands, all_reports, args.K)

            for k, cov in enumerate(curve, start=1):
                rows_curve.append(dict(app_id=app_id, variant=variant_name, k=k, report_coverage=float(cov)))

            def cov_at(k0: int) -> float:
                if k0 <= 0:
                    return 0.0
                if not curve:
                    return 0.0
                if k0 > len(curve):
                    return float(curve[-1])
                return float(curve[k0 - 1])

            rows_covsum.append(
                dict(
                    app_id=app_id,
                    variant=variant_name,
                    auc=float(auc),
                    coverage_at_10=cov_at(10),
                    coverage_at_20=cov_at(20),
                    coverage_at_50=cov_at(50),
                )
            )
            return chosen, cands, all_reports

        chosen_no, cands_no, all_reports_no = run_variant("no_unify", G_no, rep_no, type_no)
        chosen_u,  cands_u,  all_reports_u  = run_variant("unify",    G_u,  rep_u,  type_u)


        auc_no = next(r["auc"] for r in rows_covsum if r["app_id"] == app_id and r["variant"] == "no_unify")
        auc_u2 = next(r["auc"] for r in rows_covsum if r["app_id"] == app_id and r["variant"] == "unify")
        rows_delta.append(
            dict(
                app_id=app_id,
                delta_cc=int(cc_u - cc_no),
                delta_lcc_ratio=float(lcc_ratio_u - lcc_ratio_no),
                delta_cross_report_ratio=float(cer_u - cer_no),
                delta_auc=float(auc_u2 - auc_no),
            )
        )


        if args.seed_test_n and args.seed_test_n > 0:
            seeds = [s for s, _ in cands_u[: args.seed_test_n]]
            for seed in seeds:
                cov_no = candidate_subgraph_reports(G_no, seed, args.khop, rep_no, type_no, universe)
                cov_u3 = candidate_subgraph_reports(G_u,  seed, args.khop, rep_u,  type_u,  universe)
                rows_seed.append(dict(app_id=app_id, variant="no_unify", seed_entity=seed, hop=args.khop, covered_reports=len(cov_no)))
                rows_seed.append(dict(app_id=app_id, variant="unify",    seed_entity=seed, hop=args.khop, covered_reports=len(cov_u3)))


    df_graph_stats = pd.DataFrame(rows_graph_stats)
    df_conn = pd.DataFrame(rows_conn)
    df_cross = pd.DataFrame(rows_cross)
    df_curve = pd.DataFrame(rows_curve)
    df_covsum = pd.DataFrame(rows_covsum)
    df_delta = pd.DataFrame(rows_delta)
    df_seed = pd.DataFrame(rows_seed) if rows_seed else pd.DataFrame(columns=["app_id","variant","seed_entity","hop","covered_reports"])

    df_unify_stats = pd.DataFrame(rows_unify_stats) if rows_unify_stats else pd.DataFrame(
        columns=["app_id","raw_entities","unified_entities","reduction_ratio","avg_cluster_size","max_cluster_size","singleton_ratio"]
    )
    df_unify_top = pd.concat(unify_top_tables, ignore_index=True) if unify_top_tables else pd.DataFrame(
        columns=["app_id","cluster","cluster_size","example_entities"]
    )


    wb = Workbook()
    wb.remove(wb.active)

    def add_sheet(name: str, df: pd.DataFrame):
        ws = wb.create_sheet(title=name)
        write_df(ws, df)
        autosize_columns(ws)

    add_sheet("README", readme)
    add_sheet("graph_stats", df_graph_stats)
    add_sheet("connectivity_metrics", df_conn)
    add_sheet("cross_report_metrics", df_cross)
    add_sheet("coverage_curve", df_curve)
    add_sheet("coverage_summary", df_covsum)
    add_sheet("delta_metrics", df_delta)
    if df_seed.shape[0] > 0:
        add_sheet("seed_expansion", df_seed)
    if df_unify_stats.shape[0] > 0:
        add_sheet("unification_stats", df_unify_stats)
    if df_unify_top.shape[0] > 0:
        add_sheet("unification_clusters_top", df_unify_top)


    figs_ws = wb.create_sheet("FIGURES")
    figs_ws["A1"] = "Generated figures (PNGs embedded)."


    row_cursor = 3
    if df_curve.shape[0] > 0:
        for app_id in sorted(df_curve["app_id"].unique()):
            png = out_dir / f"coverage_curve_{app_id}.png"
            plot_coverage_per_app(df_curve, png, app_id)
            figs_ws[f"A{row_cursor}"] = f"Coverage Curve – {app_id}"
            img = XLImage(str(png))
            img.anchor = f"A{row_cursor+1}"
            figs_ws.add_image(img)
            row_cursor += 24

        png_all = out_dir / "coverage_curve_across_apps.png"
        plot_coverage_across_apps(df_curve, png_all)
        figs_ws[f"A{row_cursor}"] = "Coverage Curve – Across Apps (mean ± std)"
        img_all = XLImage(str(png_all))
        img_all.anchor = f"A{row_cursor+1}"
        figs_ws.add_image(img_all)

    wb.save(out_xlsx)
    print(f"[OK] Wrote: {out_xlsx}")
    print(f"[OK] Saved plots to: {out_dir}")


if __name__ == "__main__":
    main()
